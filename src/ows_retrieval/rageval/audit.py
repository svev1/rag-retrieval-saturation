"""Generate the blank manual-audit sheet: a stratified sample of answers to rate.

Draws PER_CELL answers per (retriever, k) cell (seed 42) and writes an .xlsx where
each row shows the question, the generated answer, and the documents the generator
saw, with empty Yes/Partly/No dropdown columns for faithfulness and relevancy. The
retriever/topic/k keys are written but hidden, so the rater is not anchored; they
let the filled ratings be joined back to the RAGAS scores later (see agreement.py).

This produces the *blank* sheet only. It carries no RAGAS scores by design: the
audit is done blind, and the comparison against RAGAS happens afterwards in
agreement.py. The filled sheet (with the hand ratings) is the committed artifact.
"""
import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

PER_CELL = 2
RATING = '"Yes,Partly,No"'
COLUMNS = [
    "audit_id", "question", "answer", "documents",
    "manual_faithfulness", "manual_relevancy", "notes",
    "retriever", "topic_id", "k", "n_docs",  # join keys, hidden in the sheet
]
WIDTHS = {"A": 8, "B": 38, "C": 50, "D": 95, "E": 18, "F": 18, "G": 28}
HIDDEN_COLS = ["H", "I", "J", "K"]


def sample_rows(dataset_path: Path, per_cell: int = PER_CELL, seed: int = 42) -> pd.DataFrame:
    """PER_CELL answers from each (retriever, k) cell, deterministic for reproducibility."""
    df = pd.read_parquet(dataset_path)
    return (
        df.groupby(["retriever", "k"]).sample(n=per_cell, random_state=seed)
        .sort_values(["retriever", "k", "topic_id"])
        .reset_index(drop=True)
    )


def _documents_text(contexts) -> str:
    return "\n\n".join(f"[Document {i}] {c}" for i, c in enumerate(contexts, 1))


def build_sheet(dataset_path: Path, out_path: Path, per_cell: int = PER_CELL) -> pd.DataFrame:
    sample = sample_rows(dataset_path, per_cell)
    rows = [
        {
            "audit_id": i,
            "question": r.question,
            "answer": r.answer,
            "documents": _documents_text(list(r.contexts)),
            "manual_faithfulness": "",
            "manual_relevancy": "",
            "notes": "",
            "retriever": r.retriever,
            "topic_id": int(r.topic_id),
            "k": int(r.k),
            "n_docs": int(r.n_contexts),
        }
        for i, r in enumerate(sample.itertuples(), 1)
    ]
    df = pd.DataFrame(rows, columns=COLUMNS)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(out_path, index=False, engine="openpyxl")
    _format(out_path, len(df))
    return df


def _format(path: Path, n_rows: int) -> None:
    """Wrap text, set widths, freeze the header, hide join keys, add rating dropdowns."""
    wb = load_workbook(path)
    ws = wb.active
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    for col in HIDDEN_COLS:
        ws.column_dimensions[col].hidden = True
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = wrap
    ws.freeze_panes = "B2"
    for col_letter in ("E", "F"):  # manual_faithfulness, manual_relevancy
        dv = DataValidation(type="list", formula1=RATING, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{n_rows + 1}")
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the blank manual-audit sheet (.xlsx).")
    parser.add_argument("dataset", type=Path, help="dataset.parquet")
    parser.add_argument("out", type=Path, help="output .xlsx path")
    parser.add_argument("--per-cell", type=int, default=PER_CELL, help="answers per (retriever, k) cell")
    args = parser.parse_args()

    df = build_sheet(args.dataset, args.out, args.per_cell)
    print(f"{len(df)} rows sampled ({args.per_cell} per retriever x k cell)")
    print(f"wrote {args.out} (blank — fill manual_faithfulness / manual_relevancy by hand)")


if __name__ == "__main__":
    main()
